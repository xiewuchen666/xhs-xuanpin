import csv
import io
import json
import sys
import tempfile
import unittest
from datetime import datetime, timedelta, time
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import db
import exporter
import metrics
import server


class ExporterTests(unittest.TestCase):
    def test_csv_has_utf8_bom_and_chinese_headers(self):
        data, mimetype, filename = exporter.build_export(
            "single",
            "csv",
            [{
                "title": "中文商品",
                "shop_name": "测试店铺",
                "url": "https://example.invalid/item",
                "price": 19.9,
                "total_sales": 123,
                "today": 5,
                "rolling24": 8,
                "rolling24_basis": "加入后 · 已监控8小时",
                "increment": 2,
                "interval_hours": 0.6,
                "updated_at": "2026-09-19 01:00:00",
                "status": "数据已更新",
            }],
        )
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertIn("text/csv", mimetype)
        self.assertTrue(filename.endswith(".csv"))
        text = data.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        self.assertEqual(rows[0][0], "商品标题")
        self.assertEqual(rows[1][0], "中文商品")
        self.assertEqual(rows[1][7], "加入后 · 已监控8小时")
        self.assertEqual(rows[1][9], "0.6")

    def test_xlsx_preserves_chinese_and_numeric_values(self):
        data, mimetype, filename = exporter.build_export(
            "shops",
            "xlsx",
            [{
                "shop_name": "测试店铺",
                "shop_id": "shop-1",
                "rating": "4.9",
                "brand_name": "测试品牌",
                "brand_fans_count": 12000,
                "brand_notes_count": 321,
                "product_count": 1,
                "shop_today": 5,
                "shop_today_coverage": "1/1",
                "shop_rolling24": 8,
                "shop_rolling24_coverage": "1/1",
                "product_title": "中文商品",
                "product_url": "https://example.invalid/item",
                "price": 19.9,
                "total_sales": 123,
                "today": 5,
                "rolling24": 8,
                "rolling24_basis": "完整24小时",
                "increment": 2,
                "interval_hours": 0.6,
                "updated_at": "2026-09-19 01:00:00",
                "product_status": "数据已更新",
            }],
        )
        self.assertIn("spreadsheetml", mimetype)
        self.assertTrue(filename.endswith(".xlsx"))
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        sheet = workbook["店铺监控"]
        self.assertEqual(sheet["A1"].value, "店铺名称")
        self.assertEqual(sheet["A2"].value, "测试店铺")
        self.assertEqual(sheet["L2"].value, "中文商品")
        self.assertEqual(sheet["Q2"].value, 8)
        self.assertEqual(sheet["R2"].value, "完整24小时")
        self.assertEqual(sheet["T2"].value, 0.6)


class ExportApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="xhs-xuanpin-phase4-")
        self.old_dir, self.old_path = db.DATA_DIR, db.DB_PATH
        db.DATA_DIR = Path(self.temp.name)
        db.DB_PATH = db.DATA_DIR / "phase4.db"
        db.init_db()
        self.app = server.create_app(testing=True)
        self.client = self.app.test_client()

    def tearDown(self):
        db.DATA_DIR, db.DB_PATH = self.old_dir, self.old_path
        self.temp.cleanup()

    def test_export_api_returns_attachment_for_json_and_form(self):
        rows = [{
            "title": "导出商品",
            "shop_name": "中文店铺",
            "url": "https://example.invalid/item",
            "price": 9.9,
            "total_sales": 10,
            "today": 1,
            "rolling24": 2,
            "increment": 1,
            "interval_hours": 1.0,
            "updated_at": "2026-09-19 01:00:00",
            "status": "数据已更新",
        }]
        response = self.client.post(
            "/api/export",
            json={"module": "single", "format": "csv", "rows": rows},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertTrue(response.data.startswith(b"\xef\xbb\xbf"))
        self.assertIn("导出商品", response.data.decode("utf-8-sig"))

        payload = json.dumps(
            {"module": "selection", "format": "xlsx", "rows": rows},
            ensure_ascii=False,
        )
        response = self.client.post("/api/export", data={"payload": payload})
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(response.data), data_only=True)
        self.assertEqual(workbook["选品中心"]["A2"].value, "导出商品")

    def test_export_api_rejects_invalid_format(self):
        response = self.client.post(
            "/api/export",
            json={"module": "single", "format": "pdf", "rows": []},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("仅支持", response.get_json()["error"])

    def test_sales_detail_exports_only_selected_product_and_daily_last_reading(self):
        now = metrics.now()
        yesterday = datetime.combine(now.date() - timedelta(days=1), time(12), tzinfo=metrics.TZ)

        def collect(item, when, sales):
            return db.persist(
                f"https://example.invalid/{item}",
                {"item_id": item, "title": f"商品{item}", "observed_at": metrics.text_time(when),
                 "total_sales": sales, "sales_raw": f"已售{sales}", "sales_precision": "exact"},
            )

        chosen = collect("chosen", yesterday, 10)
        collect("chosen", yesterday.replace(hour=23), 12)
        collect("chosen", now - timedelta(minutes=1), 16)
        other = collect("other", now - timedelta(minutes=1), 99)

        response = self.client.post("/api/export", json={
            "module": "single", "format": "xlsx", "detail_period": "7d", "rows": [{"id": chosen}],
        })
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        sheet = workbook["销量明细"]
        self.assertEqual([sheet.cell(4, col).value for col in range(1, 7)],
                         ["商品标题", "链接", "销量", "增量", "采集时间", "备注"])
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual((sheet["C5"].value, sheet["D5"].value), (16, 4))
        self.assertEqual(sheet["C6"].value, 12)
        self.assertEqual(sheet["A6"].value, "商品chosen")
        self.assertEqual(sheet["A5"].alignment.horizontal, "center")
        self.assertEqual(sheet["A5"].border.bottom.color.rgb, "00F2F2F2")
        self.assertTrue(sheet.tables)
        self.assertTrue(sheet.conditional_formatting)
        self.assertNotIn("商品other", [row[0].value for row in sheet.iter_rows(min_row=5)])

        rejected = self.client.post("/api/export", json={
            "module": "selection", "format": "xlsx", "detail_period": "24h", "rows": [{"id": other}],
        })
        self.assertEqual(rejected.status_code, 400)


if __name__ == "__main__":
    unittest.main()
