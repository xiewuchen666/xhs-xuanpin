from __future__ import annotations

import csv
from datetime import datetime, timedelta
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import metrics


PRODUCT_COLUMNS = [
    ("title", "商品标题"),
    ("shop_name", "店铺名称"),
    ("url", "商品链接"),
    ("price", "当前价"),
    ("total_sales", "累计销量"),
    ("today", "今日新增"),
    ("rolling24", "近24小时新增"),
    ("rolling24_basis", "近24小时统计口径"),
    ("increment", "最近区间新增"),
    ("interval_hours", "最近区间时长（小时）"),
    ("updated_at", "最近更新时间"),
    ("status", "状态"),
]

SHOP_COLUMNS = [
    ("shop_name", "店铺名称"),
    ("shop_id", "店铺ID"),
    ("rating", "店铺评分"),
    ("brand_name", "品牌/账号"),
    ("brand_fans_count", "粉丝数"),
    ("brand_notes_count", "发布笔记数"),
    ("product_count", "监控商品数"),
    ("shop_today", "店铺今日新增汇总"),
    ("shop_today_coverage", "今日新增有效覆盖"),
    ("shop_rolling24", "店铺近24小时新增汇总"),
    ("shop_rolling24_coverage", "近24小时有效覆盖"),
    ("product_title", "商品标题"),
    ("product_url", "商品链接"),
    ("price", "当前价"),
    ("total_sales", "累计销量"),
    ("today", "今日新增"),
    ("rolling24", "近24小时新增"),
    ("rolling24_basis", "近24小时统计口径"),
    ("increment", "最近区间新增"),
    ("interval_hours", "最近区间时长（小时）"),
    ("updated_at", "最近更新时间"),
    ("product_status", "商品状态"),
]

MODULE_LABELS = {
    "single": "单品监控",
    "selection": "选品中心",
    "shops": "店铺监控",
}


def _columns(module: str) -> list[tuple[str, str]]:
    if module in {"single", "selection"}:
        return PRODUCT_COLUMNS
    if module == "shops":
        return SHOP_COLUMNS
    raise ValueError("无效导出模块")


def _safe_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _normalized_rows(module: str, rows: Any) -> tuple[list[tuple[str, str]], list[list[Any]]]:
    columns = _columns(module)
    if not isinstance(rows, list):
        raise ValueError("导出数据格式错误")
    if len(rows) > 20000:
        raise ValueError("单次最多导出 20000 行")
    normalized: list[list[Any]] = []
    for item in rows:
        if not isinstance(item, dict):
            raise ValueError("导出数据格式错误")
        normalized.append([_safe_cell(item.get(key)) for key, _ in columns])
    return columns, normalized


def _filename(module: str, extension: str) -> str:
    label = MODULE_LABELS[module]
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return f"{label}-{stamp}.{extension}"


def build_csv(module: str, rows: Any) -> tuple[bytes, str, str]:
    columns, normalized = _normalized_rows(module, rows)
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow([label for _, label in columns])
    writer.writerows(normalized)
    # UTF-8 BOM makes Chinese open correctly in Windows Excel without import steps.
    data = ("\ufeff" + output.getvalue()).encode("utf-8")
    return data, "text/csv; charset=utf-8", _filename(module, "csv")


def build_xlsx(module: str, rows: Any) -> tuple[bytes, str, str]:
    columns, normalized = _normalized_rows(module, rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = MODULE_LABELS[module]

    headers = [label for _, label in columns]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in normalized:
        sheet.append(row)

    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions

    for index, (_, label) in enumerate(columns, start=1):
        values = [label]
        for row in normalized[:500]:
            value = row[index - 1]
            if value not in (None, ""):
                values.append(str(value))
        width = min(max(max((len(v) for v in values), default=len(label)) + 2, 10), 42)
        sheet.column_dimensions[get_column_letter(index)].width = width

    key_to_index = {key: idx for idx, (key, _) in enumerate(columns, start=1)}
    for key in ("price",):
        idx = key_to_index.get(key)
        if idx:
            for cell in sheet[get_column_letter(idx)][1:]:
                cell.number_format = "0.00"
    for key in ("interval_hours",):
        idx = key_to_index.get(key)
        if idx:
            for cell in sheet[get_column_letter(idx)][1:]:
                cell.number_format = "0.0"

    output = BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        _filename(module, "xlsx"),
    )


def build_export(module: str, fmt: str, rows: Any) -> tuple[bytes, str, str]:
    if fmt == "csv":
        return build_csv(module, rows)
    if fmt == "xlsx":
        return build_xlsx(module, rows)
    raise ValueError("导出格式仅支持 csv 或 xlsx")


DETAIL_PERIODS = {"24h": ("近24小时", 1), "7d": ("近7天", 7), "30d": ("近30天", 30)}
DETAIL_HEADERS = ["商品标题", "链接", "销量", "增量", "采集时间", "备注"]


def build_sales_detail(module: str, period: str, histories: list[dict[str, Any]], as_of=None) -> tuple[bytes, str, str]:
    if module not in {"single", "selection"} or period not in DETAIL_PERIODS:
        raise ValueError("无效销量明细类型")
    current = metrics.timestamp(as_of) or metrics.now()
    label, days = DETAIL_PERIODS[period]
    rows = []
    for history in histories:
        product = history["product"]
        source = sorted(
            (metrics.normalized_snapshot(item) for item in history["snapshots"]),
            key=lambda item: (item["_time"], item["id"]),
        )
        resolved = metrics.resolve_counter_rows(source)
        grouped = {}
        for item in resolved:
            key = item["collected_at"][:13 if period == "24h" else 10]
            grouped[key] = item
        keys = sorted(grouped)
        for index in range(len(keys) - 1, -1, -1):
            item = grouped[keys[index]]
            if period == "24h":
                if not current - timedelta(hours=24) <= item["_time"] <= current:
                    continue
            elif not current.date() - timedelta(days=days - 1) <= item["_time"].date() <= current.date():
                continue
            state = item.get("_counter_state")
            if state in {"pending_drop", "ignored_drop", "reset_evidence"}:
                note = "销量回落读数异常，增量不可计算"
            elif state == "reset_baseline":
                note = "销量基线重置，增量不可跨段计算"
            elif item.get("total_sales") is None or item.get("sales_precision") == "unknown":
                note = "销量读数缺失或无法解析"
            elif item.get("sales_precision") == "lower_bound":
                note = "销量为下限值，增量不可精确计算"
            elif item.get("sales_precision") == "approximate":
                note = "销量为近似值，增量不可精确计算"
            else:
                note = ""
            delta = None
            previous = grouped[keys[index - 1]] if index else None
            if previous:
                if period == "24h":
                    hours = (item["_time"] - previous["_time"]).total_seconds() / 3600
                    consecutive = keys[index - 1] == (item["_time"] - timedelta(hours=1)).strftime("%Y-%m-%d %H") and 0.5 <= hours <= 1.5
                    gap_note = "前一小时缺采或采样间隔异常，增量不可计算"
                else:
                    consecutive = previous["_time"].date() == item["_time"].date() - timedelta(days=1)
                    gap_note = "前一日缺采，增量不可计算"
                if consecutive:
                    result = metrics.interval(previous, item, resolved)
                    delta = result["value"]
                    if delta is None:
                        note = note or result["reason"]
                else:
                    note = note or gap_note
            else:
                note = note or "首次采样，无上期读数可比较"
            rows.append([
                _safe_cell(product["title"]), _safe_cell(product["url"]),
                item.get("total_sales"), delta, item["_time"].replace(tzinfo=None), _safe_cell(note),
            ])
            if len(rows) > 20000:
                raise ValueError("单次最多导出 20000 行")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销量明细"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = f"商品销量明细 · {label}"
    sheet.merge_cells("A2:F2")
    basis = "每商品每小时" if period == "24h" else "每商品每天"
    sheet["A2"] = f"{basis}取最后一次采样；销量为采集读数，增量与上期有效读数比较，无法核实时留空。"
    for column, header in enumerate(DETAIL_HEADERS, start=1):
        sheet.cell(4, column, header)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:F{sheet.max_row}"
    sheet.column_dimensions["A"].width = 58
    sheet.column_dimensions["B"].width = 57
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 16
    sheet.column_dimensions["E"].width = 25
    sheet.column_dimensions["F"].width = 58
    sheet.row_dimensions[1].height = 38
    sheet.row_dimensions[2].height = 29
    sheet.row_dimensions[4].height = 30
    line = Side(style="thin", color="F2F2F2")
    border = Border(left=line, right=line, top=line, bottom=line)
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            if cell.row == 4:
                cell.fill = PatternFill("solid", fgColor="1E3A5F")
                cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
            else:
                cell.font = Font(name="Microsoft YaHei", size=10, color="1F2937")
                if cell.column in (3, 4):
                    cell.number_format = "#,##0"
                elif cell.column == 5:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
        if row[0].row >= 5:
            sheet.row_dimensions[row[0].row].height = 23
    for address in ("A1", "A2"):
        sheet[address].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A1"].font = Font(name="Microsoft YaHei", size=16, bold=True, color="183153")
    for row in sheet["A1:F1"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="EAF2F8")
    sheet["A2"].font = Font(name="Microsoft YaHei", size=10, color="475569")
    if rows:
        table = Table(displayName="SalesDetail", ref=f"A4:F{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        sheet.add_table(table)
        sheet.conditional_formatting.add(
            f"A5:B{sheet.max_row}",
            FormulaRule(formula=["AND($A5=$A4,$B5=$B4,SUBTOTAL(103,$E4)=1)"], font=Font(color="FFFFFF")),
        )
    output = BytesIO()
    workbook.save(output)
    filename = f"{MODULE_LABELS[module]}-{label}销量明细-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename
